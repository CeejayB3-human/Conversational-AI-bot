import boto3
import json
import io
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

bedrock_agent = boto3.client('bedrock-agent-runtime', region_name='us-east-1')
bedrock = boto3.client('bedrock-runtime', region_name='us-east-1')
s3 = boto3.client('s3', region_name='us-east-1')

KB_ID    = os.environ.get('KB_ID', '2QKASOKWYP')
MODEL_ID = os.environ.get('MODEL_ID', 'arn:aws:bedrock:us-east-1:047719615900:inference-profile/global.anthropic.claude-sonnet-4-6')
BUCKET   = os.environ.get('BUCKET', 'smarthome-ai-bucket')

SYSTEM_PROMPT = """You are Soji, a friendly and knowledgeable smart home advisor for Smart Lagos Homes.

Smart Lagos Homes sells and installs smart home technology for Nigerian households.
You help customers understand and choose the right smart home products for their needs.

Your personality:
- Warm, friendly and conversational — like a knowledgeable friend
- Speak plain English — avoid technical jargon
- Relate to Nigerian home challenges: NEPA outages, security concerns, high electricity bills
- Be budget-conscious and practical

Your job:
- Ask about the customer's home (size, location in Lagos, number of bedrooms)
- Ask about their main pain points (power outages, security, cooling, water, high bills)
- Ask about their budget
- Recommend the right Smart Lagos Homes products and packages
- Explain in plain language why each product helps them
- Give prices in Nigerian Naira (NGN)

IMPORTANT — PROPOSAL GENERATION RULE:
After your THIRD response to the customer, you MUST end your message with [GENERATE_PROPOSAL]
on its own line — no exceptions. By the third response you have enough information.
Stop asking questions and make your final recommendation.
Do not tell the customer you are generating a document. Just include [GENERATE_PROPOSAL] at the very end.

Rules:
- Only recommend products from the Smart Lagos Homes catalog
- Always follow the installation priority: Power first, then Security, then Cooling, then Water, then Lighting
- Be honest about costs — do not invent prices
- Never ask for an email address
- Keep responses concise — maximum 4 short paragraphs
- Do not use markdown formatting like ** or # in your responses
- Write in plain conversational text only"""

PROPOSAL_PROMPT = """You are a senior sales consultant at Smart Lagos Homes.

Based on the conversation below, generate a complete smart home proposal.

Include these sections in plain text:

CUSTOMER SUMMARY
Home type, location, bedrooms, main pain points

RECOMMENDED SMART HOME PACKAGE
List each product with part number, plain description and price in NGN

INSTALLATION PRIORITY
Step by step order with brief reason

TOTAL COST BREAKDOWN
Each item with price, subtotal, discount if applicable, grand total in NGN

ESTIMATED MONTHLY SAVINGS
Realistic estimate of monthly electricity savings

WHAT CHANGES IN YOUR HOME
Plain language daily life improvements

NEXT STEPS
Call to action

Write in warm professional language. Plain text only.
No markdown, no asterisks, no hashtags.
CAPITALISED TEXT for section headings only."""

NAVY = RGBColor(0x06, 0x0d, 0x1f)
GOLD = RGBColor(0xe8, 0xb8, 0x30)
MUTED = RGBColor(0x55, 0x55, 0x55)

def build_docx(proposal_text, ref_id):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    h = doc.add_paragraph()
    r = h.add_run('SMART LAGOS HOMES')
    r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY

    s = doc.add_paragraph()
    rs = s.add_run('Smart Home Advisor — Customer Proposal')
    rs.font.size = Pt(12); rs.font.color.rgb = GOLD

    m = doc.add_paragraph()
    rm = m.add_run(f'Reference: {ref_id}  |  Date: {datetime.now().strftime("%d %B %Y")}  |  Valid: 30 days')
    rm.font.size = Pt(9); rm.font.color.rgb = MUTED
    m.paragraph_format.space_after = Pt(16)

    HEADERS = ['CUSTOMER SUMMARY','RECOMMENDED','INSTALLATION','TOTAL COST','ESTIMATED','WHAT CHANGES','NEXT STEPS']

    def is_hdr(line):
        return any(line.upper().strip().startswith(h) for h in HEADERS)

    for line in proposal_text.split('\n'):
        clean = line.strip()
        if not clean:
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if is_hdr(clean):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(clean.upper())
            r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NAVY
        elif clean.startswith('-'):
            p = doc.add_paragraph(style='List Bullet')
            r = p.add_run(clean.lstrip('-').strip())
            r.font.size = Pt(11)
        else:
            p = doc.add_paragraph()
            r = p.add_run(clean)
            r.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(3)

    doc.add_paragraph()
    f = doc.add_paragraph()
    rf = f.add_run('Smart Lagos Homes  |  www.smartlagoshomes.com.ng  |  hello@smartlagoshomes.com.ng  |  0801 234 5678')
    rf.font.size = Pt(9); rf.font.color.rgb = MUTED
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

def generate_proposal(history, ref_id):
    conversation = ''
    for h in history:
        role = 'Soji' if h.get('role') == 'assistant' else 'Customer'
        conversation += f"{role}: {h.get('content','')}\n\n"

    last_msg = history[-1].get('content','') if history else ''
    retrieval = bedrock_agent.retrieve(
        knowledgeBaseId=KB_ID,
        retrievalQuery={'text': last_msg[:500]},
        retrievalConfiguration={'vectorSearchConfiguration': {'numberOfResults': 5}}
    )
    catalog = ''
    for r in retrieval.get('retrievalResults',[]):
        catalog += r.get('content',{}).get('text','') + '\n\n'

    resp = bedrock.invoke_model(
        modelId=MODEL_ID,
        body=json.dumps({
            'anthropic_version': 'bedrock-2023-05-31',
            'max_tokens': 3000,
            'system': PROPOSAL_PROMPT,
            'messages': [{'role': 'user', 'content': f"Catalog:\n{catalog}\n\nConversation:\n{conversation}\n\nGenerate proposal. Reference: {ref_id}"}]
        })
    )
    text = json.loads(resp['body'].read())['content'][0]['text']
    buf = build_docx(text, ref_id)

    key = f'proposals/{ref_id}-smart-home-proposal.docx'
    s3.put_object(Bucket=BUCKET, Key=key, Body=buf.read(),
        ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

    url = s3.generate_presigned_url('get_object',
        Params={'Bucket': BUCKET, 'Key': key}, ExpiresIn=604800)
    return url

def lambda_handler(event, context):
    body = event.get('body', '{}')
    if isinstance(body, str):
        body = json.loads(body)

    message = body.get('message', '')
    history = body.get('history', [])

    headers = {
        'Access-Control-Allow-Origin': '*',
        'Access-Control-Allow-Headers': 'Content-Type',
        'Access-Control-Allow-Methods': 'POST, OPTIONS'
    }

    if not message:
        return {'statusCode': 400, 'headers': headers, 'body': json.dumps({'error': 'Message required'})}

    retrieval = bedrock_agent.retrieve(
        knowledgeBaseId=KB_ID,
        retrievalQuery={'text': message},
        retrievalConfiguration={'vectorSearchConfiguration': {'numberOfResults': 3}}
    )
    catalog = ''
    for r in retrieval.get('retrievalResults', []):
        catalog += r.get('content', {}).get('text', '') + '\n\n'

    messages = []
    for h in history:
        messages.append({'role': h.get('role'), 'content': h.get('content')})
    messages.append({'role': 'user', 'content': f"Catalog info:\n{catalog}\n\nCustomer: {message}"})

    resp = bedrock.invoke_model(
        modelId=MODEL_ID,
        body=json.dumps({
            'anthropic_version': 'bedrock-2023-05-31',
            'max_tokens': 600,
            'system': SYSTEM_PROMPT,
            'messages': messages
        })
    )
    reply = json.loads(resp['body'].read())['content'][0]['text']

    proposal_url = None
    clean_reply = reply

    if '[GENERATE_PROPOSAL]' in reply:
        clean_reply = reply.replace('[GENERATE_PROPOSAL]', '').strip()
        ref_id = 'SLH-' + datetime.now().strftime('%Y%m%d%H%M%S')
        try:
            full_history = history + [
                {'role': 'user', 'content': message},
                {'role': 'assistant', 'content': clean_reply}
            ]
            proposal_url = generate_proposal(full_history, ref_id)
        except Exception as e:
            print(f"Proposal generation error: {e}")

    return {
        'statusCode': 200,
        'headers': headers,
        'body': json.dumps({
            'reply': clean_reply,
            'proposal_url': proposal_url,
            'status': 'success'
        })
    }
